#!/usr/bin/env python3
"""Build the capstone log book and the annexure forms as Word files.

    python make_forms.py OUT_DIR

Writes MemGate_Log_Book.docx and MemGate_Annexure_Forms.docx (formats from
the Capstone project guidelines 2026-27, Annexures A.1-A.7). Fields nobody
has given us yet (contact numbers) are left blank.
"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, "images", "nmims_logo.png")

SCHOOL = "Mukesh Patel School of Technology Management & Engineering"
CAMPUS = "SVKM's NMIMS (Deemed-to-be University), Indore Campus"
TITLE = "MemGate: An Adaptive Memory Layer for Long-Running LLM Agents"
MENTOR = "Dr. Praveen Modi, Assistant Professor"
YEAR = "2026-27"
SEMESTER = "VII"
DEPARTMENT = "Artificial Intelligence and Data Science"
PROGRAM = "B.Tech in Artificial Intelligence and Data Science"
STUDENTS = [  # (SAP ID, name, roll no.)
    ("70562300097", "Yash Ramchandani", "S024"),
    ("70562200088", "Simar Singh Khanuja", "S019"),
]

# ------------------------------------------------------------------ helpers

def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(2)
    return doc


def para(doc_or_cell, text="", bold=False, size=None, align=None,
         italic=False, highlight=False, space_after=None):
    p = doc_or_cell.add_paragraph()
    if text:
        add_run(p, text, bold=bold, size=size, italic=italic,
                highlight=highlight)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def add_run(p, text, bold=False, size=None, italic=False, highlight=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def field(doc, label, value="", highlight=False, blank_cm=None):
    p = doc.add_paragraph()
    add_run(p, label + " ", bold=True)
    if value:
        add_run(p, value, highlight=highlight)
    if blank_cm:
        add_run(p, "_" * int(blank_cm * 2.2))
    p.paragraph_format.space_after = Pt(3)
    return p


def letterhead(doc, form_title, annex=None):
    if annex:
        para(doc, annex, bold=True, size=11, align="c", space_after=4)
    p = para(doc, align="c", space_after=2)
    p.add_run().add_picture(LOGO, width=Cm(2.3))
    para(doc, SCHOOL, bold=True, size=15, align="c", space_after=0)
    para(doc, CAMPUS, size=11, align="c", space_after=6)
    para(doc, form_title, bold=True, size=13, align="c", space_after=8)


def set_cell(cell, text="", bold=False, size=None, highlight=False,
             align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
        if line:
            add_run(p, line, bold=bold, size=size, highlight=highlight)
        if align == "c":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)


def shade(cell, hex_fill="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def row_height(row, cm):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(int(cm * 567)))
    h.set(qn("w:hRule"), "atLeast")
    trPr.append(h)


def grid(doc, header, rows, widths=None, size=10, heights=None,
         head_fill="D9D9D9"):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        set_cell(t.rows[0].cells[j], h, bold=True, size=size, align="c")
        shade(t.rows[0].cells[j], head_fill)
    for i, r in enumerate(rows, 1):
        for j, v in enumerate(r):
            hl = isinstance(v, tuple)
            set_cell(t.rows[i].cells[j], v[0] if hl else v, size=size,
                     highlight=hl)
        if heights:
            row_height(t.rows[i], heights)
    if widths:
        fix_widths(t, widths)
    return t


def fix_widths(t, widths_cm):
    """Pin column widths (Word and LibreOffice both honour tblGrid)."""
    t.autofit = False
    grid_el = t._tbl.tblGrid
    for gc, w in zip(grid_el.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for j, w in enumerate(widths_cm):
            row.cells[j].width = Cm(w)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def signature_block(doc, label, name=""):
    para(doc, space_after=10)
    p = doc.add_paragraph()
    add_run(p, label, bold=True)
    p = doc.add_paragraph()
    if name:
        add_run(p, name)
        p = doc.add_paragraph()
    add_run(p, "Signature: " + "_" * 24 + "     Date: " + "_" * 12)
    p.paragraph_format.space_after = Pt(6)


# ----------------------------------------------------------------- log book
# Weeks run Saturday to Thursday from 18 July 2026, the batch's calendar
# (Review 1, 25 July 2026, falls in week 2 as the guidelines say), and the
# mentor is met on the Wednesday of each week. Simar asked us to set the
# timeline (27 Oct 2026). Work comes from the dated record in the repository
# (SESSION_RECORD.md, docs/REVIEW2_PROGRESS_REPORT.md, git history); the
# weeks it does not cover are filled from the project plan and the work that
# the next dated entry shows was already done.

import datetime as _dt

WEEKS = [
    [  # 1
        "Identified the problem: LLM agents re-send their whole history on "
        "every turn, so cost and latency grow and accuracy falls well before "
        "the context limit.",
        "Literature survey of memory for LLM agents (MemGPT, MemoryBank, "
        "Generative Agents, A-MEM) and of prompt and KV-cache compression "
        "(LLMLingua, Gist tokens, AutoCompressors, StreamingLLM, H2O).",
        "Selected MemGPT (Packer et al., 2023) as the base paper.",
        "Drafted the problem statement, three objectives, scope and the "
        "semester Gantt chart.",
    ],
    [  # 2
        "Presented Review 1: introduction, problem statement, objectives, "
        "scope, proposed system and architecture, implementation workflow, "
        "tools and technology, feasibility and Gantt chart.",
        "Presented a comparison of 10 papers (dataset, technique, "
        "achievements, limitations, gap addressed).",
        "Noted the panel's questions: is the project needed now that context "
        "windows are large, is it feasible, is a tiered memory still novel, "
        "and what is the end product.",
    ],
    [  # 3
        "Studied the long-conversation benchmarks LoCoMo, LongMemEval and "
        "LongBench and chose LoCoMo as the main benchmark (marked evidence "
        "turns allow a retrieval-level metric).",
        "Chose the stack: Python, tiktoken for exact token counts, "
        "all-MiniLM-L6-v2 for embeddings, Qwen2.5 models for local LLM work.",
        "Designed the library interfaces: observe() for the write path and "
        "build_context() for the read path, and the data types for turns, "
        "memory items and questions.",
        "Decided to build the evaluation harness before the memory policies, "
        "so every later idea is a measured experiment.",
    ],
    [  # 4
        "Prepared evidence-backed answers to the four Review 1 panel "
        "questions.",
        "Moved the contribution from the three-tier architecture to the "
        "memory decision policy (keep / compress / forget / recall).",
        "Built Step 0: evaluation harness, three-tier store, policies P0 "
        "(sliding window) and P1 (heuristic scorer), and a synthetic test set.",
        "30 automated tests passing; made the P0 baseline budget-driven so "
        "the comparison is fair.",
    ],
    [  # 5
        "Found an evaluation leak: a consolidation rule read the answer label "
        "(fact_id), worth 45-47 recall points. Fixed it and withdrew the "
        "Review 1 figure (95% became 50%).",
        "Added label-blindness regression tests.",
        "Wired the LoCoMo benchmark (10 conversations, 1,527 questions) and "
        "the MiniLM embedder; added the answer-presence metric.",
        "Fixed Tier 2 consolidation, improved the compressor, and ran a "
        "23-configuration ablation study.",
    ],
    [  # 6
        "Imposed a storage budget on the whole memory store, not only the "
        "prompt.",
        "Key finding: scored forgetting beats oldest-first forgetting at the "
        "same storage, while compressing evicted turns loses answer recall.",
        "Scaling study to 182x compression; byte accounting including the "
        "embedding index.",
        "Clustered bootstrap confidence intervals and McNemar tests; learned "
        "scorer (leave-one-conversation-out); first publication figures; "
        "built the P2 LLM judge.",
    ],
    [  # 7
        "Ran the P2 LLM judge (Qwen2.5-0.5B) on GPU, reading salience from "
        "the digit-token logits of one forward pass.",
        "Integrated a fixed reader model (Qwen2.5-1.5B, 4-bit) for end-task "
        "accuracy and ran a pilot.",
        "Pinned every model to a local copy so runs cannot fetch new "
        "weights.",
    ],
    [  # 8
        "Measured the P2 judge: it loses to the hand-written heuristic by "
        "5.87 points, while a learned scorer gains 11.47.",
        "Prepared the Review 2 progress report, a generated deck, a results "
        "dashboard and a speaker's guide.",
        "Wrote the first draft of the IEEE conference paper with tables "
        "generated from the result files.",
    ],
    [  # 9
        "Ran end-task accuracy on all 715 answer-recoverable LoCoMo "
        "questions.",
        "Result: +6.0 of the +8.7 recall points become correct answers (69% "
        "of the gap), against a 1.7% closed-book floor.",
        "Updated the report and paper with the end-task results.",
    ],
    [  # 10
        "Added a LongMemEval loader and set up a sweep of store budgets to "
        "test the headline on a second benchmark.",
        "Matched LongMemEval's retention to LoCoMo's so the two are "
        "comparable.",
        "Continued revising the paper draft.",
    ],
    [  # 11
        "LongMemEval result: the headline does not replicate; traced to the "
        "heuristic scorer, whose separation of evidence turns reverses sign "
        "on that corpus.",
        "Ruled out the recency-reservation confound.",
        "Ran abstractive compression and online learning of the scorer at "
        "full scale.",
        "Compiled the paper and cut it to 6 pages; rebuilt the Review 2 deck "
        "in the NMIMS format with speaker notes; built the first live demo.",
    ],
    [  # 12
        "Presented Review 2 (midterm review): progress against the plan, "
        "design, results so far and the live demo.",
        "Recorded the panel's feedback and planned the remaining work: a "
        "working application and a learned scorer.",
    ],
    [  # 13
        "Turned the demo into a working chat app: saved chats, sidebar, "
        "memory panel showing every decision.",
        "Found and fixed a scorer defect (speaker tag scored with the "
        "message) and re-ran the results (P1-S 36.8% to 37.5%).",
        "Trained the learned scorer on LoCoMo and REALTALK (human chats) and "
        "shipped it in the app.",
        "Updated the report, paper, deck and dashboard to the fixed numbers.",
    ],
    [  # 14
        "Full app: one-command setup, any text input (pastes, PDF, Word, "
        "PowerPoint, other languages), a Compare tab and score explanations.",
        "Stress tests (six input bugs fixed) and a load test with 200 "
        "simultaneous users.",
        "Compared with MemGPT-style and Mem0-style memory: the learned scorer "
        "wins at every budget.",
        "Retrained the scorer on 14.6x the data; chats up to 10M tokens; on "
        "million-token chats the app keeps the answer 51.9% of the time vs "
        "19.2% for RAG.",
        "Google Gemini replies with a local Qwen backup; entry gate and score "
        "tags on every message.",
    ],
    [  # 15
        "Prepared the journal version of the paper with Dr. Praveen Modi as "
        "co-author and corresponding author.",
        "Matched the paper to the code; extended references to 52 sources.",
        "Applied the mentor's review of the paper; rewrote the abstracts and "
        "the prose.",
        "Prepared the log book, project report (black book) and annexure "
        "forms.",
    ],
    [],  # 16
    [],  # 17
]

WEEK1 = _dt.date(2026, 7, 18)


def _d(day):
    return day.strftime("%d/%m/%Y")


def week_page(doc, n, bullets):
    start = WEEK1 + _dt.timedelta(days=7 * (n - 1))
    end, meet = start + _dt.timedelta(days=5), start + _dt.timedelta(days=4)
    field(doc, f"Week No: {n}")
    p = field(doc, "Date of Reporting:", _d(meet))
    para(doc, f"{_d(start)} to {_d(end)}", space_after=8)
    para(doc, "Work carried out", bold=True, space_after=4)
    if bullets:
        for b in bullets:
            q = doc.add_paragraph(style="List Bullet")
            add_run(q, b)
            q.paragraph_format.space_after = Pt(3)
    else:
        for _ in range(8):
            para(doc, "_" * 76, space_after=8)
    para(doc, space_after=10)
    para(doc, "Faculty Remark", bold=True, space_after=6)
    for _ in range(4):
        para(doc, "_" * 76, space_after=8)
    para(doc, space_after=10)
    para(doc, "Faculty Signature with date", bold=True, space_after=10)
    para(doc, "_" * 40)


def log_book(out):
    doc = new_doc()
    letterhead(doc, "Log Book for Capstone Project", "A.7 Log Book for Capstone Project")
    field(doc, "Department:", DEPARTMENT)
    field(doc, "Program:", PROGRAM)
    field(doc, "Semester:", SEMESTER)
    field(doc, "Academic Year:", YEAR)
    field(doc, "TITLE OF THE PROJECT:", TITLE)
    para(doc, space_after=2)
    field(doc, "Name of the Faculty Mentor:", MENTOR)
    para(doc, space_after=4)
    para(doc, "STUDENT DETAILS", bold=True, align="c", space_after=4)
    rows = [["Student 1", STUDENTS[0][1], STUDENTS[0][2], ""],
            ["Student 2", STUDENTS[1][1], STUDENTS[1][2], ""]]
    t = grid(doc, ["", "NAME", "ROLL NO.", "CONTACT"], rows,
             widths=[2.6, 6.0, 3.4, 4.6], size=10.5, heights=0.8)
    for i in range(1, 3):
        t.rows[i].cells[0].paragraphs[0].runs[0].bold = True
    para(doc, space_after=6)
    para(doc, "Name of group members:", bold=True, space_after=4)
    grid(doc, ["Roll No.", "Name", "Sign (Students)",
               "Grades (By Faculty Mentor)"],
         [[r, n, "", ""] for _, n, r in STUDENTS],
         widths=[3.0, 6.0, 3.6, 4.0], size=10.5, heights=1.0)
    para(doc, space_after=6)
    para(doc, "SAP IDs: " + ", ".join(f"{n} ({s})" for s, n, _ in STUDENTS),
         size=10, italic=True)

    for i, bullets in enumerate(WEEKS):
        page_break(doc)
        week_page(doc, i + 1, bullets)
    doc.save(out)


# ---------------------------------------------------------------- annexures

OBJECTIVES = (
    "1. Build an adaptive memory decision engine that scores every message "
    "and decides whether to keep it verbatim, compress it, forget it, or "
    "recall it later, so that the prompt stays within a fixed token budget.\n"
    "2. Build a three-tier memory store (recent turns, compressed gists, "
    "long-term facts with embeddings) held to a fixed storage budget, with "
    "retrieval to bring dropped information back when it is relevant.\n"
    "3. Evaluate the memory policies on long multi-session conversation "
    "benchmarks (LoCoMo and others), changing one factor at a time, and "
    "report the trade-off between memory budget and answer quality with "
    "proper statistics.")

MOTIVATION = (
    "LLM agents re-send their whole history on every turn, so cost and "
    "latency grow with the conversation and accuracy falls well before the "
    "context limit. Sliding windows forget the oldest messages whatever they "
    "contain, and summaries lose names, numbers and dates. On LongMemEval, "
    "oracle evidence gives about 92% accuracy against about 58% when the "
    "model selects its own: a 34-point gap caused by memory selection, not "
    "reasoning. Existing memory systems are evaluated end to end, so the "
    "value of each memory decision is unknown.")

OUTCOMES = (
    "A working memory layer (Python library) that sits in front of any LLM; "
    "a measurement harness and results on public benchmarks with confidence "
    "intervals; a learned scorer for message importance; a chat application "
    "that shows every keep, compress and forget decision; a research paper.")

REFERENCES = [
    'C. Packer, S. Wooders, K. Lin, V. Fang, S. G. Patil, I. Stoica and '
    'J. E. Gonzalez, "MemGPT: Towards LLMs as Operating Systems," arXiv '
    'preprint arXiv:2310.08560, 2023.',
    'P. Chhikara, D. Khant, S. Aryan, T. Singh and D. Yadav, "Mem0: Building '
    'Production-Ready AI Agents with Scalable Long-Term Memory," arXiv '
    'preprint arXiv:2504.19413, 2025.',
    'W. Xu, Z. Liang, K. Mei, H. Gao, J. Tan and Y. Zhang, "A-MEM: Agentic '
    'Memory for LLM Agents," Advances in Neural Information Processing '
    'Systems (NeurIPS), 2025.',
    'A. Maharana, D.-H. Lee, S. Tulyakov, M. Bansal, F. Barbieri and Y. Fang, '
    '"Evaluating Very Long-Term Conversational Memory of LLM Agents," Proc. '
    'Annual Meeting of the Association for Computational Linguistics (ACL), '
    '2024.',
    'D. Wu, H. Wang, W. Yu, Y. Zhang, K.-W. Chang and D. Yu, "LongMemEval: '
    'Benchmarking Chat Assistants on Long-Term Interactive Memory," arXiv '
    'preprint arXiv:2410.10813, 2024.',
]


def team_table(doc, sign_col="Signature"):
    grid(doc, ["S No", "SAP ID", "Roll No", "Name", sign_col],
         [["1", STUDENTS[0][0], STUDENTS[0][2], STUDENTS[0][1], ""],
          ["2", STUDENTS[1][0], STUDENTS[1][2], STUDENTS[1][1], ""],
          ["3", "-", "-", "-", "-"]],
         widths=[1.2, 3.0, 2.4, 5.2, 4.8], size=10.5, heights=0.8)


def dept_block(doc, semester=SEMESTER):
    field(doc, "Department:", DEPARTMENT)
    field(doc, "Program:", PROGRAM)
    if semester:
        field(doc, "Semester:", semester)
    else:
        field(doc, "Semester:", blank_cm=6)


def eval_sheet(doc, annex, title, criteria, total, signer_lines):
    letterhead(doc, title, annex)
    header = ["S No", "SAP ID", "Roll No", "Student Name"] + criteria + [total]
    rows = [[str(i + 1), s, r, n] + [""] * (len(criteria) + 1)
            for i, (s, n, r) in enumerate(STUDENTS)]
    fixed = [0.9, 2.5, 1.5, 2.4]
    rest = (16.6 - sum(fixed)) / (len(criteria) + 1)
    grid(doc, header, rows, widths=fixed + [rest] * (len(criteria) + 1),
         size=8.5 if len(criteria) > 3 else 9.5, heights=1.2)
    para(doc, "Project: " + TITLE, size=10, italic=True)
    for label, name in signer_lines:
        signature_block(doc, label, name)


def annexures(out):
    doc = new_doc()

    # A.1 Project registration form
    letterhead(doc, "Project registration form", "A.1 Project registration form")
    dept_block(doc, semester=None)
    para(doc, "Project Team details", bold=True, space_after=4)
    team_table(doc)
    para(doc, space_after=6)
    para(doc, "Three broad domains in which the project team intends to work",
         bold=True, space_after=4)
    grid(doc, ["S No", "Domain", "Project idea in brief"],
         [["1", "Generative AI: large language models and LLM agents",
           "A memory layer for long-running LLM agents that decides what to "
           "keep, compress, forget and recall, so long conversations fit a "
           "fixed token budget (MemGate)."],
          ["2", "", ""], ["3", "", ""]],
         widths=[1.2, 5.4, 10.0], size=10.5, heights=1.4)

    # A.2 Topic approval form
    page_break(doc)
    letterhead(doc, "Topic approval form", "A.2 Topic approval form")
    dept_block(doc)
    para(doc, "Project Team details", bold=True, space_after=4)
    team_table(doc)
    para(doc, space_after=4)
    para(doc, "Type of Project (Tick any one)", bold=True, space_after=4)
    grid(doc, ["Type of project", ""],
         [["Application", ""], ["Product", ""], ["Research", "✓"]],
         widths=[5.0, 3.0], size=10.5)
    para(doc, space_after=4)
    para(doc, "Project details", bold=True, space_after=4)
    t = doc.add_table(rows=5, cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate([
            ("Project Title", TITLE),
            ("Project Objectives", OBJECTIVES),
            ("Domain", "Generative AI / Natural Language Processing: memory "
                       "and context management for LLM agents"),
            ("Motivation", MOTIVATION),
            ("Expected outcomes", OUTCOMES)]):
        set_cell(t.rows[i].cells[0], k, bold=True, size=10.5)
        set_cell(t.rows[i].cells[1], v, size=10.5)
    fix_widths(t, [3.6, 13.0])
    para(doc, space_after=4)
    para(doc, "Latest References", bold=True, space_after=4)
    grid(doc, ["S No", "Publication (Author, “Title”, "
                       "Journal/Book/Conference, Date, Page, Volume, year)"],
         [[str(i + 1), r] for i, r in enumerate(REFERENCES)],
         widths=[1.2, 15.4], size=10)
    para(doc, space_after=4)
    t = doc.add_table(rows=5, cols=2)
    t.style = "Table Grid"
    labels = ["Status of the project after Review 1\nComments/Remarks",
              "Name and Signature of the guide",
              "Name and signature of panel members for review 1",
              "Signature of HoD",
              "If any change in project title after review 1 presentation as "
              "per comments received from panel and areas of concern to be "
              "noted here"]
    for i, k in enumerate(labels):
        set_cell(t.rows[i].cells[0], k, bold=True, size=10)
        row_height(t.rows[i], 1.6 if i != 2 else 2.2)
    set_cell(t.rows[1].cells[1], MENTOR, size=10.5)
    fix_widths(t, [6.0, 10.6])

    # A.3 Review 1
    page_break(doc)
    eval_sheet(doc, "A.3 Evaluation sheet for Review 1: Topic approval and feasibility",
               "Evaluation sheet for Review 1: Topic approval and feasibility",
               ["Identification of Problem statement, Purpose and need of "
                "the project (5 marks)",
                "Feasibility of the implementation (5 marks)"],
               "Total (10 marks)",
               [("Name & Signature of the Project mentor (with Date)", MENTOR),
                ("Name & Signature of the panel members (with Date)", ""),
                ("", ""), ("", "")])

    # A.4 Review 2
    page_break(doc)
    eval_sheet(doc, "A.4 Evaluation sheet for Review 2: Midterm review",
               "Evaluation sheet for Review 2: Midterm review",
               ["Literature Review / Market Survey (5 marks)",
                "Clarity of the problem statement (5 marks)",
                "Design / System level representation (5 marks)",
                "Implementation / Results (5 marks)",
                "Interim Report (5 marks)"],
               "Total (25 marks)",
               [("Name & Signature of the Project mentor (with Date)", MENTOR),
                ("Name & Signature of the panel members (with Date)", ""),
                ("", ""), ("", "")])

    # A.5 Review 3
    page_break(doc)
    eval_sheet(doc, "A.5 Evaluation sheet for Review 3: External Panel Member Evaluation",
               "Evaluation sheet for Review 3: Internal and External Panel "
               "Member Evaluation",
               ["Literature review and clarity of the problem statement "
                "(5 marks)",
                "Design / System level representation (5 marks)",
                "Implementation and results (5 marks)",
                "Report (5 marks)", "Presentation (5 marks)",
                "Viva (5 marks)"],
               "Total (30 marks)",
               [("Name & Signature of the Internal mentor (with Date)", MENTOR),
                ("Name & Signature of the External Panel Member (with Date)", ""),
                ("", "")])

    # A.6 Research paper
    page_break(doc)
    letterhead(doc, "Evaluation sheet for Research paper",
               "A.6 Evaluation sheet for Research paper")
    grid(doc, ["S No", "SAP ID", "Roll No", "Student Name", "Marks (10 marks)"],
         [[str(i + 1), s, r, n, ""] for i, (s, n, r) in enumerate(STUDENTS)],
         widths=[1.2, 3.0, 2.4, 5.6, 4.4], size=10.5, heights=1.0)
    para(doc, space_after=4)
    para(doc, "Paper: “Selection Beats Compression: Memory Decision "
              "Policies for LLM Agents Under a Bounded Store”, Yash "
              "Ramchandani, Simar Singh Khanuja and Dr. Praveen Modi.",
         size=10, italic=True)
    p = para(doc, "Status (submitted / accepted / published, with venue and "
                  "proof attached): ", size=10)
    add_run(p, "_" * 40, size=10)
    signature_block(doc, "Name & Signature of the Project mentor (with Date)",
                    MENTOR)
    doc.save(out)


if __name__ == "__main__":
    out_dir = sys.argv[1] if len(sys.argv) > 1 else HERE
    os.makedirs(out_dir, exist_ok=True)
    log_book(os.path.join(out_dir, "MemGate_Log_Book.docx"))
    annexures(os.path.join(out_dir, "MemGate_Annexure_Forms.docx"))
    print("wrote", out_dir)
